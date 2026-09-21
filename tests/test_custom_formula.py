"""Analyst-defined formulae: the whitelist that keeps them safe, the
config coherence rules, and the non-standard marking that follows a run
using one into every export.

The security property is the one that matters most here and it is tested
adversarially: the evaluator must reject every construct that could reach
outside arithmetic, and it must do so by having no branch for them rather
than by listing them, so a construct nobody thought of is refused too.
"""

from __future__ import annotations

import numpy as np
import pandas as pd
import pytest
from pydantic import ValidationError

from pricelab.core.config import IndexConfig, RunConfig
from pricelab.core.security import FormulaError, validate_formula_syntax
from pricelab.engine import custom
from pricelab.engine.index import build_index


def _panel():
    periods = pd.date_range("2020-01-01", periods=4, freq="MS")
    return pd.DataFrame({
        "period": list(periods) * 2,
        "item_id": ["1"] * 4 + ["2"] * 4,
        "price_imputed": [10.0, 11.0, 12.0, 13.0, 20.0, 19.0, 22.0, 26.0],
    }), periods


# ---------------------------------------------------------------------
# It must never be eval
# ---------------------------------------------------------------------
@pytest.mark.parametrize("hostile", [
    "__import__('os').system('echo pwned')",
    "os.system('echo pwned')",
    "().__class__.__bases__[0].__subclasses__()",
    "open('/etc/passwd').read()",
    "(lambda: 1)()",
    "[x for x in range(10)]",
    "{'a': 1}['a']",
    "jevons if jevons else dutot",
    "globals()",
    "jevons.__class__",
    "exec('x=1')",
])
def test_hostile_expressions_are_refused(hostile):
    """Every one of these must be refused. They are refused because
    `_eval_node` has no branch for the construct, not because the
    construct appears on a denylist -- which is what makes the same
    refusal apply to constructs this test did not think of."""
    with pytest.raises(FormulaError):
        validate_formula_syntax(hostile, custom.ELEMENTARY_VARIABLES)


def test_a_hostile_expression_is_refused_at_config_time_not_run_time():
    """The rejection happens when the formula is written, so a compile
    cannot fail half way through after the quality and imputation stages
    have already run."""
    with pytest.raises(ValidationError, match="not a permitted expression"):
        IndexConfig(formula="custom", custom_formula="__import__('os').system('x')")


def test_an_unknown_name_is_refused_with_the_available_names_listed():
    with pytest.raises(ValidationError, match="unknown name"):
        IndexConfig(formula="custom", custom_formula="jevons + inflation_target")


# ---------------------------------------------------------------------
# Config coherence
# ---------------------------------------------------------------------
def test_custom_without_an_expression_is_refused():
    with pytest.raises(ValidationError, match="no expression"):
        IndexConfig(formula="custom")


def test_an_expression_that_would_be_ignored_is_refused():
    """A custom_formula set alongside a standard formula name is
    ambiguous about which the caller meant, and the ambiguity is resolved
    by refusing rather than by silently preferring one."""
    with pytest.raises(ValidationError, match="would be silently ignored"):
        IndexConfig(formula="jevons", custom_formula="cswd")


def test_a_valid_custom_formula_is_accepted():
    cfg = IndexConfig(formula="custom", custom_formula="0.5 * jevons + 0.5 * dutot")
    assert cfg.custom_formula == "0.5 * jevons + 0.5 * dutot"


# ---------------------------------------------------------------------
# It computes the right numbers
# ---------------------------------------------------------------------
def test_a_custom_formula_naming_one_building_block_reproduces_it_exactly():
    df, _periods = _panel()
    custom_result = build_index(df, IndexConfig(formula="custom", custom_formula="jevons"))
    standard = build_index(df, IndexConfig(formula="jevons"))
    pd.testing.assert_series_equal(custom_result["index"], standard["index"])


def test_a_custom_formula_can_express_cswd_from_its_two_biased_components():
    """The example that shows the feature is worth having: an analyst can
    build CSWD out of Carli and the harmonic mean without this tool
    having to ship it as a named option."""
    df, _periods = _panel()
    built = build_index(df, IndexConfig(
        formula="custom", custom_formula="(carli * harmonic_mean) ** 0.5"))
    shipped = build_index(df, IndexConfig(formula="custom", custom_formula="cswd"))
    pd.testing.assert_series_equal(built["index"], shipped["index"])


def test_the_pandas_and_numpy_paths_agree():
    """`engine.custom.evaluate_elementary` (pandas, for a single
    comparison) and `engine.index._custom_np` (numpy, for the period
    loop) must produce the same number, or a custom formula would mean
    one thing when inspected and another when compiled."""
    items = ["1", "2", "3"]
    a = pd.Series([10.0, 20.0, 5.0], index=items)
    b = pd.Series([11.0, 19.0, 6.0], index=items)
    expression = "0.5 * jevons + 0.5 * dutot"

    from pricelab.engine.index import _custom_np

    assert custom.evaluate_elementary(expression, a, b).value == pytest.approx(
        _custom_np(a.to_numpy(), b.to_numpy(), expression))


def test_an_elementary_result_carries_the_expression_that_produced_it():
    items = ["1", "2"]
    a = pd.Series([10.0, 20.0], index=items)
    b = pd.Series([11.0, 19.0], index=items)
    result = custom.evaluate_elementary("cswd", a, b)

    assert result.formula == "custom"
    assert result.parameters["expression"] == "cswd"
    assert "jevons" in result.parameters["inputs"]
    assert result.n_items == 2


def test_an_elementary_custom_result_is_nan_with_nothing_to_compare():
    empty = pd.Series(dtype=float)
    assert np.isnan(custom.evaluate_elementary("jevons", empty, empty).value)


# ---------------------------------------------------------------------
# Aggregate-level custom formulae
# ---------------------------------------------------------------------
def _index_frame():
    periods = pd.date_range("2020-01-01", periods=3, freq="MS")
    return pd.DataFrame({
        "Bread": [100.0, 105.0, 110.0],
        "Milk": [100.0, 98.0, 96.0],
        "All items": [100.0, 101.5, 102.9],
    }, index=periods)


def test_an_aggregate_formula_can_exclude_a_category():
    """The core-inflation case the restricted evaluator was built for:
    an index defined as the headline with one volatile category taken
    out."""
    I = _index_frame()
    series = custom.evaluate_aggregate("2 * all_items - bread", I)
    assert series.tolist() == pytest.approx([100.0, 98.0, 95.8])


def test_aggregate_variables_sanitise_category_labels():
    I = pd.DataFrame({"Food & drink": [100.0], "All items": [100.0]},
                     index=pd.date_range("2020-01-01", periods=1, freq="MS"))
    names = custom.aggregate_variables(I.columns)
    assert names["food_drink"] == "Food & drink"
    assert names["all_items"] == "All items"


def test_colliding_category_names_are_refused_rather_than_resolved():
    """Two labels reducing to the same variable name would make any
    formula using it ambiguous, and an index whose definition is
    ambiguous is worse than one that failed to compile."""
    I = pd.DataFrame({"Food & drink": [100.0], "Food/drink": [100.0]},
                     index=pd.date_range("2020-01-01", periods=1, freq="MS"))
    with pytest.raises(FormulaError, match="both reduce to"):
        custom.aggregate_variables(I.columns)


def test_an_aggregate_formula_yields_nan_for_a_period_with_a_missing_input():
    """A category that starts late is an ordinary fact about a
    collection; the periods that are defined are still produced."""
    I = _index_frame()
    I.loc[I.index[1], "Bread"] = np.nan
    series = custom.evaluate_aggregate("bread + milk", I)
    assert np.isnan(series.iloc[1])
    assert not np.isnan(series.iloc[0])


def test_validate_aggregate_formula_checks_against_the_categories_present():
    I = _index_frame()
    assert custom.validate_aggregate_formula("bread + milk", I.columns) == {"bread", "milk"}
    with pytest.raises(FormulaError, match="unknown name"):
        custom.validate_aggregate_formula("bread + cheese", I.columns)


def test_validate_elementary_formula_returns_the_names_referenced():
    assert custom.validate_elementary_formula("0.5*jevons + 0.5*dutot") == {"jevons", "dutot"}


# ---------------------------------------------------------------------
# Non-standard marking
# ---------------------------------------------------------------------
def test_a_standard_run_is_not_marked():
    cfg = IndexConfig(formula="jevons")
    assert not custom.is_non_standard(cfg)
    assert custom.non_standard_notice(cfg) == ""
    assert custom.custom_formula_parameters(cfg) == {}


def test_a_custom_run_is_marked_and_the_notice_names_the_expression():
    cfg = IndexConfig(formula="custom", custom_formula="0.5*jevons + 0.5*dutot")
    assert custom.is_non_standard(cfg)
    notice = custom.non_standard_notice(cfg)
    assert "NON-STANDARD INDEX" in notice
    assert "0.5*jevons + 0.5*dutot" in notice
    assert "not comparable" in notice
    assert custom.custom_formula_parameters(cfg)["expression"] == "elementary: 0.5*jevons + 0.5*dutot"


def test_a_custom_aggregate_formula_also_marks_the_run_non_standard():
    """The aggregate escape hatch (Phase 3 Task 4's second half, wired to
    the interface in Phase 10.5) is marked exactly like the elementary one,
    and a run using both names both."""
    cfg = IndexConfig(custom_aggregate_formula="(bread + milk) / 2")
    assert cfg.formula == "jevons" and custom.is_non_standard(cfg)
    assert "aggregate: (bread + milk) / 2" in custom.non_standard_notice(cfg)
    both = IndexConfig(formula="custom", custom_formula="cswd",
                       custom_aggregate_formula="all_items")
    assert custom.non_standard_expression(both) == "elementary: cswd; aggregate: all_items"
    with pytest.raises(ValueError, match="not a permitted expression"):
        IndexConfig(custom_aggregate_formula="__import__('os')")
    with pytest.raises(ValueError, match="empty"):
        IndexConfig(custom_aggregate_formula="   ")


def test_the_custom_formula_is_recorded_in_the_config_like_any_other_parameter():
    """It rides in the config JSON, so everything that already covers a
    parameter covers it: the registry's content hash, the cache key, and
    the saved configuration of every run."""
    import json

    cfg = RunConfig(index=IndexConfig(formula="custom", custom_formula="cswd"))
    restored = RunConfig.from_dict(json.loads(cfg.to_json()))
    assert restored.index.formula == "custom"
    assert restored.index.custom_formula == "cswd"


def test_two_runs_differing_only_in_their_custom_formula_hash_differently():
    from pricelab.core.cache import content_key

    a = IndexConfig(formula="custom", custom_formula="jevons")
    b = IndexConfig(formula="custom", custom_formula="dutot")
    assert content_key(b"same file", a.model_dump_json(), "label") != \
        content_key(b"same file", b.model_dump_json(), "label")


def test_the_method_note_names_the_expression_rather_than_the_word_custom():
    """"Custom" on its own tells a reader nothing they can check, so the
    written method note quotes the actual expression."""
    from pricelab import run_pipeline
    from pricelab.reporting.report import method_note

    periods = pd.date_range("2020-01-01", periods=3, freq="MS")
    df = pd.DataFrame({
        "period": list(periods) * 2,
        "category": ["Bread"] * 3 + ["Milk"] * 3,
        "item_id": ["1"] * 3 + ["2"] * 3,
        "item_name": ["White"] * 3 + ["Semi"] * 3,
        "price_reported": [1.0, 1.1, 1.2, 2.0, 1.9, 2.2],
    })
    cfg = RunConfig(index=IndexConfig(formula="custom", custom_formula="cswd"))
    res = run_pipeline(df, cfg)

    note = method_note(res)
    assert "NON-STANDARD INDEX" in note
    assert "analyst-defined (cswd)" in note


def test_every_export_surface_marks_a_non_standard_run():
    """The deck, the written report and a CSV download all carry the
    mark. A caveat that appears in one export and not another is a caveat
    that stops travelling as soon as someone forwards the wrong file."""
    from io import BytesIO

    from pptx import Presentation

    from pricelab import build_all_charts, build_deck, run_pipeline
    from pricelab.core.security import safe_csv_with_notice
    from pricelab.engine.insights import build_narrative
    from pricelab.reporting.report import build_markdown

    periods = pd.date_range("2020-01-01", periods=4, freq="MS")
    df = pd.DataFrame({
        "period": list(periods) * 2,
        "category": ["Bread"] * 4 + ["Milk"] * 4,
        "item_id": ["1"] * 4 + ["2"] * 4,
        "item_name": ["White"] * 4 + ["Semi"] * 4,
        "price_reported": [1.0, 1.1, 1.2, 1.3, 2.0, 1.9, 2.2, 2.6],
    })
    cfg = RunConfig(index=IndexConfig(formula="custom", custom_formula="cswd"))
    res = run_pipeline(df, cfg)
    nar = build_narrative(res)
    charts = build_all_charts(res)

    deck_text = "\n".join(
        shape.text_frame.text
        for slide in Presentation(BytesIO(build_deck(res, nar, charts, "custom run"))).slides
        for shape in slide.shapes if shape.has_text_frame)
    assert "NON-STANDARD INDEX" in deck_text

    assert "NON-STANDARD INDEX" in build_markdown(res, nar, "custom run")

    csv = safe_csv_with_notice(res["indices"], custom.non_standard_notice(cfg.index))
    assert csv.splitlines()[0].startswith("# NON-STANDARD INDEX")


def test_the_csv_notice_is_sanitised_like_any_other_exported_text():
    """The notice quotes a user-supplied expression, so it is not exempt
    from the formula-injection rule just because the application wrote
    the sentence around it."""
    from pricelab.core.security import safe_csv_with_notice

    df = pd.DataFrame({"a": [1]})
    out = safe_csv_with_notice(df, "=cmd|'/c calc'!A1")
    assert out.splitlines()[0].startswith("# '=")
