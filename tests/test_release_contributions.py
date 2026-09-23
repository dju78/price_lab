"""Phase 7b, Task 0a: contributions in the release outputs.

The bulletin, the Word report and the deck each carry what drove the
headline: every category's contribution, the level of the tree the
contributions are at, and the residual between their sum and the published
change -- shown as a row of the table, not absorbed. A run without weights
gets the reason instead of a silently missing section.
"""

from __future__ import annotations

import io
import re

import numpy as np
import pandas as pd
import pytest
from dbtarget import database_url

from pricelab import (
    build_all_charts,
    build_deck,
    build_docx,
    build_markdown,
    build_narrative,
    infer_schema,
    run_pipeline,
    standardise,
)
from pricelab.core import db
from pricelab.core.config import RunConfig, get_settings
from pricelab.core.provenance import build_stamp
from pricelab.reporting.report import CONTRIBUTIONS_TITLE, contributions_summary


def _collection(weighted: bool, months: int = 26) -> pd.DataFrame:
    rng = np.random.default_rng(1)
    rows = []
    for category, weight, drift in (("Bread", 40.0, 0.004), ("Fuel", 25.0, 0.009),
                                    ("Rent", 35.0, 0.002)):
        for item in range(4):
            price = 10.0 + item
            for t, period in enumerate(pd.date_range("2023-01-01", periods=months, freq="MS")):
                if t:
                    price *= 1 + drift + rng.normal(0, 0.003)
                row = {"Date": period, "Category": category, "Item_ID": f"{category}{item}",
                       "Item_Name": f"{category} {item}", "Reported_Price": round(price, 4)}
                if weighted:
                    row["Weight"] = weight / 4
                rows.append(row)
    raw = pd.DataFrame(rows)
    return standardise(raw, infer_schema(raw))


@pytest.fixture(scope="module")
def weighted_run() -> tuple[pd.DataFrame, dict]:
    df = _collection(weighted=True)
    return df, run_pipeline(df, RunConfig(label="weighted release"))


def test_the_summary_states_the_level_and_shows_the_residual(weighted_run):
    _, res = weighted_run
    note, table = contributions_summary(res)
    assert "level 1 of the classification tree" in note
    assert "on the same period a year earlier" in note
    assert "shown in the table rather than absorbed" in note
    categories = table.drop(index=[i for i in table.index if str(i) not in {"Bread", "Fuel", "Rent"}])
    headline = res["indices"]["All items"]
    published = (headline.iloc[-1] / headline.iloc[-13] - 1) * 100
    assert categories["Contribution, pp"].sum() == pytest.approx(published, abs=1e-5)
    residual = table.loc["Residual, pp (sum minus published change)", "Contribution, pp"]
    assert abs(residual) < 1e-10
    assert table.loc["All items change, % (as published)", "Contribution, pp"] == pytest.approx(
        published, abs=1e-6)


def test_a_residual_that_is_not_zero_is_shown_not_hidden(weighted_run):
    """Replace the headline with something the categories do not add up to
    (as an analyst-defined aggregate formula would) and the residual row
    shows the gap rather than any category silently absorbing it."""
    _, res = weighted_run
    tampered = dict(res)
    indices = res["indices"].copy()
    indices["All items"] = indices["All items"] * np.linspace(1.0, 1.01, len(indices))
    tampered["indices"] = indices
    _, table = contributions_summary(tampered)
    residual = table.loc["Residual, pp (sum minus published change)", "Contribution, pp"]
    assert abs(residual) > 0.1


def test_a_run_without_weights_says_why_there_are_no_contributions():
    res = run_pipeline(_collection(weighted=False), RunConfig())
    note, table = contributions_summary(res)
    assert table.empty
    assert "no expenditure weights" in note and "geometric mean" in note


def _docx_text(data: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(data))
    cells = [c.text for t in document.tables for row in t.rows for c in row.cells]
    return "\n".join([p.text for p in document.paragraphs] + cells)


def _pptx_text(data: bytes) -> str:
    from pptx import Presentation

    return "\n".join("\n".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)
                     for slide in Presentation(io.BytesIO(data)).slides)


def test_the_word_report_markdown_and_deck_carry_the_contributions(weighted_run):
    _, res = weighted_run
    nar = build_narrative(res)
    charts = build_all_charts(res)
    stamp = build_stamp(res, "weighted release")
    for name, text in (
            ("word", _docx_text(build_docx(res, nar, charts, "weighted release", stamp))),
            ("markdown", build_markdown(res, nar, "weighted release")),
            ("deck", _pptx_text(build_deck(res, nar, charts, "weighted release")))):
        flat = re.sub(r"\s+", " ", text)
        assert CONTRIBUTIONS_TITLE in flat, name
        assert "level 1 of the classification tree" in flat, name
        assert "Residual" in flat, name
        for category in ("Bread", "Fuel", "Rent"):
            assert category in flat, name


def test_the_bulletin_carries_the_contributions(weighted_run, tmp_path, monkeypatch):
    from pypdf import PdfReader

    from pricelab.core.registry import register_run
    from pricelab.reporting import bulletin

    df, res = weighted_run
    monkeypatch.setenv("PRICELAB_DATABASE_URL", database_url(tmp_path, "release.db"))
    monkeypatch.setenv("PRICELAB_STORE_DIR", str(tmp_path / "store"))
    get_settings.cache_clear()
    db.reset_db_state()
    db.init_db()
    try:
        with db.session_scope() as s:
            run = register_run(s, df, res["config"], "weighted release", result=res)
            stamp = build_stamp(res, "weighted release", run=run)
            pdf = bulletin.build_bulletin(res, build_narrative(res), build_all_charts(res),
                                          stamp, run=run)
    finally:
        db.reset_db_state()
        get_settings.cache_clear()
    text = re.sub(r"\s+", " ", "\n".join(page.extract_text() or ""
                                         for page in PdfReader(io.BytesIO(pdf)).pages))
    assert CONTRIBUTIONS_TITLE in text
    assert "level 1 of the classification tree" in text
    assert "Residual, pp" in text
    assert "Sum of contributions" in text


def test_the_excel_pack_carries_the_contributions_on_their_own_sheet(weighted_run):
    """Phase 8, Task 0a: the evidence pack, where an auditor checks the
    arithmetic, has the table with its level and its residual."""
    from openpyxl import load_workbook

    from pricelab.reporting.excel import SHEETS, build_evidence_pack

    df, res = weighted_run
    book = load_workbook(io.BytesIO(build_evidence_pack(res, build_stamp(res, "w"), source=df)))
    assert "Contributions" in SHEETS and "Contributions" in book.sheetnames
    cells = [[c for c in row] for row in book["Contributions"].iter_rows(values_only=True)]
    text = " ".join(str(c) for row in cells for c in row if c is not None)
    assert "level 1 of the classification tree" in text
    labels = [row[0] for row in cells if row and row[0]]
    assert "Residual, pp (sum minus published change)" in labels
    assert {"Bread", "Fuel", "Rent", "Sum of contributions"} <= set(labels)
    rows = {row[0]: row for row in cells if row and row[0]}
    total = sum(rows[c][3] for c in ("Bread", "Fuel", "Rent"))
    assert total == pytest.approx(rows["Sum of contributions"][3], abs=1e-5)


def test_an_unweighted_run_s_pack_says_why_there_are_no_contributions():
    from openpyxl import load_workbook

    from pricelab.reporting.excel import build_evidence_pack

    df = _collection(weighted=False)
    res = run_pipeline(df, RunConfig())
    book = load_workbook(io.BytesIO(build_evidence_pack(res, build_stamp(res, "u"), source=df)))
    text = " ".join(str(c) for row in book["Contributions"].iter_rows(values_only=True)
                    for c in row if c is not None)
    assert "no expenditure weights" in text
