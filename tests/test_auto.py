"""Tests for automatic schema inference, self-configuration and export.

These cover the promise the product makes: upload a file, get a correct
analysis and usable outputs without being asked anything.
"""

import io
import numpy as np
import pandas as pd
import pytest

from pricelab import (infer_schema, auto_configure, analyse, standardise,
                      build_deck, build_docx, build_markdown, method_note)


def collection(n_items=4, n_periods=36, seed=0, category="Bread"):
    rng = np.random.default_rng(seed)
    periods = pd.date_range("2020-01-01", periods=n_periods, freq="MS")
    rows = []
    for i in range(n_items):
        p = 1.5 + 0.2 * i
        for t in periods:
            rows.append({"Date": t, "Category": category, "Item_ID": f"{category[:2]}{i}",
                         "Item_Name": f"{category} item {i}", "Reported_Price": round(p, 4)})
            p *= 1.003 * (1 + rng.normal(0, 0.004))
    return pd.DataFrame(rows)


def two_categories(seed=0):
    return pd.concat([collection(seed=seed, category="Bread"),
                      collection(seed=seed + 1, category="Pasta")], ignore_index=True)


# ----------------------------------------------------------------------
def test_schema_inferred_from_conventional_names():
    s = infer_schema(collection())
    assert (s.date, s.item_id, s.category, s.price) == \
           ("Date", "Item_ID", "Category", "Reported_Price")


def test_schema_inferred_from_unconventional_names():
    df = collection().rename(columns={"Date": "month", "Reported_Price": "value_gbp",
                                      "Item_ID": "sku", "Category": "product_group",
                                      "Item_Name": "description"})
    s = infer_schema(df)
    assert s.date == "month" and s.price == "value_gbp"
    assert s.item_id == "sku" and s.category == "product_group"


def test_category_and_item_not_swapped():
    """Category must be the low-cardinality column even when names are ambiguous."""
    df = two_categories().rename(columns={"Category": "id_group", "Item_ID": "id_item"})
    s = infer_schema(df)
    assert df[s.category].nunique() < df[s.item_id].nunique()


# ----------------------------------------------------------------------
def test_seasonal_gap_gets_seasonal_hold():
    df = two_categories()
    off = df["Date"].dt.month.isin([11, 12, 1, 2]) & (df["Category"] == "Pasta")
    df.loc[off, "Reported_Price"] = 0
    cfg, decisions = auto_configure(standardise(df))
    assert cfg.imputation.by_category["Pasta"] == "seasonal_hold"
    assert any("seasonal" in d.lower() for d in decisions)


def test_collection_gap_gets_class_mean():
    df = two_categories()
    gap = df["Date"].isin(pd.to_datetime(["2020-06-01", "2020-07-01"])) & \
        (df["Category"] == "Bread")
    df.loc[gap, "Reported_Price"] = 0
    cfg, _ = auto_configure(standardise(df))
    assert cfg.imputation.by_category["Bread"] == "class_mean"


def test_reference_window_follows_periodicity():
    monthly = standardise(collection(n_periods=36))
    assert auto_configure(monthly)[0].quality.reference_window == 13
    q = collection(n_periods=16)
    q["Date"] = pd.date_range("2020-01-01", periods=16, freq="QS").tolist() * 4
    assert auto_configure(standardise(q))[0].quality.reference_window == 5


def test_weights_switch_the_formula():
    df = two_categories()
    df["Weight"] = 1.0
    cfg, decisions = auto_configure(standardise(df, infer_schema(df)))
    assert cfg.index.formula == "laspeyres"
    assert any("weight" in d.lower() for d in decisions)


def test_every_decision_is_explained():
    """A decision the user cannot see is one they cannot defend."""
    df = two_categories()
    df.loc[df["Date"].dt.month.isin([12, 1]) & (df["Category"] == "Pasta"),
           "Reported_Price"] = 0
    cfg, decisions = auto_configure(standardise(df))
    assert len(decisions) >= 3
    assert all(isinstance(d, str) and len(d) > 40 for d in decisions)


# ----------------------------------------------------------------------
def test_analyse_end_to_end_without_configuration():
    out = analyse(standardise(two_categories()), "Test collection")
    assert out["narrative"] is not None
    assert len(out["narrative"].findings) >= 3
    assert "index" in out["charts"]
    assert out["narrative"].headline


def test_findings_are_ranked():
    out = analyse(standardise(two_categories()), "Test")
    scores = [f.importance for f in out["narrative"].findings]
    assert scores == sorted(scores, reverse=True)


def test_every_finding_carries_evidence():
    out = analyse(standardise(two_categories()), "Test")
    for f in out["narrative"].findings:
        assert f.headline and f.detail
        assert f.evidence or f.table is not None


# ----------------------------------------------------------------------
def test_deck_is_a_valid_pptx():
    from pptx import Presentation
    out = analyse(standardise(two_categories()), "Test")
    data = build_deck(out["result"], out["narrative"], dict(out["charts"]), "Test")
    prs = Presentation(io.BytesIO(data))
    assert len(prs.slides) >= 5


def test_deck_slides_carry_speaker_notes():
    """Every slide carries speaker notes, not just the finding slides."""
    from pptx import Presentation
    out = analyse(standardise(two_categories()), "Test")
    prs = Presentation(io.BytesIO(
        build_deck(out["result"], out["narrative"], dict(out["charts"]), "Test")))
    for i, s in enumerate(prs.slides):
        assert s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip(), \
            f"slide {i} has no speaker notes"


def test_report_is_a_valid_docx_with_method_note():
    from docx import Document
    out = analyse(standardise(two_categories()), "Test")
    doc = Document(io.BytesIO(
        build_docx(out["result"], out["narrative"], dict(out["charts"]), "Test")))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Method note" in text
    assert "Limitations" in text


def test_markdown_report_contains_the_headline():
    out = analyse(standardise(two_categories()), "Test")
    md = build_markdown(out["result"], out["narrative"], "Test")
    assert out["narrative"].headline in md
    assert "Method note" in md


def test_method_note_states_the_limitations():
    """The note must not overclaim. Quality adjustment and weights are absent
    and the note has to say so."""
    out = analyse(standardise(two_categories()), "Test")
    note = method_note(out["result"])
    assert "quality adjustment" in note.lower()
    assert "weight" in note.lower()


# ----------------------------------------------------------------------
def test_handles_a_single_category():
    out = analyse(standardise(collection()), "One category")
    assert out["narrative"] is not None


def test_handles_a_collection_with_no_faults():
    """A clean file must produce a clean report, not invent problems."""
    out = analyse(standardise(two_categories()), "Clean")
    q = out["result"]["quality"]["flag_summary"]["count"]
    assert int(q.get("scale_error_x100", 0)) == 0
    assert not any("unit error" in f.headline for f in out["narrative"].findings)
