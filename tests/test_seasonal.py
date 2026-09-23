"""Phase 6, Task 1: strictly seasonal items, the Rothwell index,
counter-seasonal estimation and seasonal adjustment.

The two rules this phase enforces are what most of these tests are about.
Whichever engine actually ran must be named in every output -- so there are
tests for the X-13 path, for the STL fallback, and for the labelling of
each, including on a machine (this one) where the X-13 binary is absent and
the successful path has to be exercised through a substituted runner. And
the unadjusted series must appear alongside the adjusted one everywhere, so
there are tests that go looking for it in the publication table, the CSV,
the Markdown report, the Word report and the Excel pack rather than trusting
that each format remembered.
"""

from __future__ import annotations

import io
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
import streamlit
from dbtarget import database_url
from streamlit.testing.v1 import AppTest

from pricelab import build_markdown, build_narrative, infer_schema, run_pipeline, standardise
from pricelab.core import db
from pricelab.core.config import RunConfig, SeasonalConfig, get_settings
from pricelab.engine import seasonal as sn

REPO_ROOT = Path(__file__).resolve().parents[1]
PRICES = REPO_ROOT / "supermarket_price_collection.xlsx"

# The bundled collection is the only one with a real season and real faults in
# it, so most of this module is about that file. Skipped rather than failed
# where it is absent, matching tests/test_phase3_hard_gate.py -- though the
# Dockerfile copies it in precisely so these do run in the image.
pytestmark = pytest.mark.skipif(not PRICES.exists(),
                                reason="fixture workbook not present")


# ---------------------------------------------------------------------
# Panels
# ---------------------------------------------------------------------
def _panel(n_years: int = 4, *, seasonal_months: tuple[int, ...] = (5, 6, 7, 8),
           n_regular: int = 4, n_seasonal: int = 2, trend: float = 0.002,
           amplitude: float = 0.0, start: str = "2020-01-01") -> pd.DataFrame:
    """A monthly panel with regular items and strictly seasonal ones.

    `amplitude` puts a genuine seasonal wave into the *regular* items'
    prices -- which is what seasonal adjustment is for, and a different
    thing from an item being off the shelf, which is what `seasonal_months`
    controls. Keeping the two separable in the fixture is what lets the
    tests below tell them apart.
    """
    periods = pd.date_range(start, periods=12 * n_years, freq="MS")
    rows: list[dict[str, object]] = []
    for t, period in enumerate(periods):
        wave = 1.0 + amplitude * np.sin(2 * np.pi * (period.month - 1) / 12)
        for i in range(n_regular):
            rows.append({"period": period, "category": "Staples", "item_id": f"R{i}",
                         "price_imputed": (10.0 + i) * (1 + trend) ** t * wave,
                         "price_clean": (10.0 + i) * (1 + trend) ** t * wave,
                         "weight": 1.0, "imputation": ""})
        if period.month in seasonal_months:
            for i in range(n_seasonal):
                rows.append({"period": period, "category": "Berries", "item_id": f"S{i}",
                             "price_imputed": (4.0 + i) * (1 + trend) ** t,
                             "price_clean": (4.0 + i) * (1 + trend) ** t,
                             "weight": 1.0, "imputation": ""})
        # One item in the seasonal category that is on the shelf all year,
        # so the category has something to carry it under confinement.
        rows.append({"period": period, "category": "Berries", "item_id": "S_ALLYEAR",
                     "price_imputed": 6.0 * (1 + trend) ** t,
                     "price_clean": 6.0 * (1 + trend) ** t, "weight": 1.0, "imputation": ""})
    return pd.DataFrame(rows)


@pytest.fixture(scope="module")
def collection() -> pd.DataFrame:
    raw = pd.read_excel(PRICES)
    return standardise(raw, infer_schema(raw))


@pytest.fixture(scope="module")
def compiled(collection: pd.DataFrame) -> dict:
    return run_pipeline(collection, RunConfig())


# ---------------------------------------------------------------------
# Strictly seasonal items
# ---------------------------------------------------------------------
def test_an_item_off_the_shelf_the_same_months_every_year_is_strictly_seasonal():
    items = sn.strictly_seasonal_items(_panel())
    seasonal = items[items["strictly_seasonal"]]
    assert set(seasonal.index) == {"S0", "S1"}
    assert seasonal.loc["S0", "in_season"] == (5, 6, 7, 8)
    assert seasonal.loc["S0", "out_of_season"] == (1, 2, 3, 4, 9, 10, 11, 12)
    assert seasonal.loc["S0", "years_observed"] == 4


def test_an_item_merely_interrupted_is_not_seasonal_and_the_reason_is_given():
    """A season is a repetition. Applying a seasonal rule to a collection
    failure invents a pattern and then removes it, so the distinction is
    enforced and the count that failed is reported rather than left as a
    bare no."""
    frame = _panel(n_years=3)
    gap = (frame["item_id"] == "R0") & frame["period"].isin(
        pd.date_range("2021-03-01", periods=3, freq="MS"))
    items = sn.strictly_seasonal_items(frame[~gap])
    assert not items.loc["R0", "strictly_seasonal"]
    # Those calendar months are still observed in the other two years, so
    # the item has no off *season* at all -- which is the point: one gap,
    # however long, is not a repetition.
    assert "priced in every calendar period" in items.loc["R0", "excluded_because"]

    # An item missing the same single month every year does repeat, but is
    # on the shelf 11 months of 12, above the ceiling that separates a
    # season from an item with gaps.
    december = (frame["item_id"] == "R1") & (frame["period"].dt.month == 12)
    items = sn.strictly_seasonal_items(frame[~december])
    assert not items.loc["R1", "strictly_seasonal"]
    assert "gaps, not a season" in items.loc["R1", "excluded_because"]


def test_one_year_of_data_cannot_establish_a_season():
    items = sn.strictly_seasonal_items(_panel(n_years=1))
    assert not items["strictly_seasonal"].any()
    assert "fewer than the 2 needed" in items.loc["S0", "excluded_because"]


def test_the_supermarket_collection_finds_its_seasonal_items(compiled):
    items = sn.strictly_seasonal_items(compiled["imputed"])
    qualifying = items[items["strictly_seasonal"]]
    assert len(qualifying) >= 3
    # Every one of them is off the shelf in winter and on it in summer,
    # which is the fixture's strawberry season.
    for out_of_season in qualifying["out_of_season"]:
        assert 1 in out_of_season and 2 in out_of_season
    for in_season in qualifying["in_season"]:
        assert 6 in in_season and 7 in in_season


def test_a_frame_without_the_columns_a_seasonal_analysis_needs_is_refused():
    with pytest.raises(sn.SeasonalError, match="period"):
        sn.strictly_seasonal_items(pd.DataFrame({"item_id": ["A"], "price_imputed": [1.0]}))


# ---------------------------------------------------------------------
# The two treatments
# ---------------------------------------------------------------------
def test_class_confinement_and_weight_update_differ_on_the_same_data():
    """The acceptance criterion. Both are compiled, both are reported, and
    the gap between them is stated in the units the headline is published
    in -- because the choice between them is a judgement about what a basket
    means while part of it is off the shelf."""
    comparison = sn.compare_treatments(_panel())
    assert set(comparison.results) == {"class_confinement", "weight_update"}
    assert not comparison.identical
    assert comparison.max_gap_pp > 0.01
    assert np.isfinite(comparison.final_gap_pp)
    assert {"class_confinement", "weight_update", "gap_pp"} <= set(comparison.table.columns)
    assert comparison.seasonal_items == ("S0", "S1")


def test_the_two_treatments_agree_exactly_when_nothing_is_strictly_seasonal():
    """The control that makes the previous test mean something. With no
    seasonal item, no weight ever leaves the basket, the two treatments hand
    the same aggregator the same weights, and the series are identical to
    machine precision. A gap here would prove the comparison was measuring
    the aggregation rather than the seasonality."""
    frame = _panel(n_years=4, seasonal_months=tuple(range(1, 13)))
    comparison = sn.compare_treatments(frame)
    assert not comparison.seasonal_items
    assert comparison.identical
    assert comparison.max_gap_pp == pytest.approx(0.0, abs=1e-9)
    assert any("necessarily agree" in n for n in comparison.notes)


def test_the_two_treatments_agree_when_the_seasonal_items_are_removed(compiled):
    """The same control on the real collection: strip the strictly seasonal
    items and the two must coincide exactly."""
    frame = compiled["imputed"]
    items = sn.strictly_seasonal_items(frame)
    without = frame[~frame["item_id"].isin(items.index[items["strictly_seasonal"]])]
    assert sn.compare_treatments(without, compiled["config"].index).max_gap_pp == \
        pytest.approx(0.0, abs=1e-9)
    assert sn.compare_treatments(frame, compiled["config"].index).max_gap_pp > 0.01


def test_confinement_keeps_a_class_weight_constant_and_weight_update_does_not():
    comparison = sn.compare_treatments(_panel())
    confined = comparison.results["class_confinement"].weights_by_period["Berries"]
    updated = comparison.results["weight_update"].weights_by_period["Berries"]
    assert confined.nunique() == 1
    assert updated.nunique() > 1
    # In season the two coincide; out of season the seasonal items' weight
    # has left the basket under weight update and not under confinement.
    in_season = updated.index.month.isin([5, 6, 7, 8])
    assert updated[in_season].to_numpy() == pytest.approx(confined[in_season].to_numpy())
    assert (updated[~in_season] < confined[~in_season]).all()


def test_a_period_in_which_no_price_moves_leaves_the_weight_update_level_where_it_was():
    """The arithmetic property that makes weighting movements rather than
    levels correct under weights that change. A weighted mean of levels
    under moving weights drifts when the weights move and no price does,
    which would read as inflation."""
    frame = _panel(n_years=3, trend=0.0)
    result = sn.weight_update(frame)
    levels = result.headline.dropna()
    assert levels.to_numpy() == pytest.approx(np.full(len(levels), 100.0))


# ---------------------------------------------------------------------
# The Rothwell index
# ---------------------------------------------------------------------
def test_the_rothwell_index_prices_each_period_against_base_year_averages():
    frame = _panel(n_years=3, trend=0.0)
    result = sn.rothwell(frame, base_year=2020)
    assert result.base_year == 2020
    assert result.items_in_base == 4 + 2 + 1
    # With no trend and no within-year price movement, every period prices
    # its own items at exactly their base-year averages.
    assert result.index.to_numpy() == pytest.approx(
        np.full(len(result.index), 100.0), rel=1e-9)
    # The item count moves with the season, which is the composition effect
    # inherent to the form and the reason it is reported.
    assert result.items_by_period.nunique() > 1


def test_the_rothwell_index_follows_a_proportional_price_rise():
    frame = _panel(n_years=3, trend=0.0)
    later = frame["period"] >= pd.Timestamp("2021-01-01")
    frame.loc[later, ["price_imputed", "price_clean"]] *= 1.10
    result = sn.rothwell(frame, base_year=2020)
    assert float(result.index.loc[pd.Timestamp("2021-06-01")]) == pytest.approx(110.0, rel=1e-9)


def test_a_rothwell_index_refuses_a_base_year_the_collection_does_not_have():
    with pytest.raises(sn.SeasonalError, match="base year 1999"):
        sn.rothwell(_panel(n_years=2), base_year=1999)


def test_the_rothwell_index_says_which_items_its_base_year_never_saw():
    frame = _panel(n_years=3)
    newcomer = frame[frame["item_id"] == "R0"].copy()
    newcomer["item_id"] = "NEW"
    newcomer = newcomer[newcomer["period"] >= pd.Timestamp("2022-01-01")]
    result = sn.rothwell(pd.concat([frame, newcomer], ignore_index=True), base_year=2020)
    assert any("never appear in base year 2020" in n for n in result.notes)


# ---------------------------------------------------------------------
# Counter-seasonal estimation
# ---------------------------------------------------------------------
def test_counter_seasonal_estimates_move_with_the_items_that_are_in_season():
    """Not held flat. A flat off-season price says the item did not move
    while everything around it did, and delivers the whole accumulated
    difference in one month when it returns."""
    frame = _panel(n_years=3, trend=0.01)
    result = sn.counter_seasonal_estimate(frame)
    assert result.items == ("S0", "S1")
    assert result.estimated > 0

    filled = result.frame
    estimates = filled[filled["imputation"] == "counter_seasonal"]
    assert set(estimates["item_id"]) <= {"S0", "S1"}
    # Every estimate is strictly above the last in-season price it was
    # carried from, because the in-season items rose through the gap.
    s0 = filled[filled["item_id"] == "S0"].sort_values("period")
    september = s0[s0["period"] == pd.Timestamp("2020-09-01")]["price_imputed"].iloc[0]
    august = s0[s0["period"] == pd.Timestamp("2020-08-01")]["price_imputed"].iloc[0]
    assert september > august


def test_counter_seasonal_estimates_are_labelled_as_constructions():
    result = sn.counter_seasonal_estimate(_panel(n_years=3))
    estimates = result.frame[result.frame["imputation"] == "counter_seasonal"]
    assert len(estimates) == result.estimated
    assert (result.coverage["share"] > 0).any()
    assert any("construction, not an observation" in n for n in result.notes)


def test_counter_seasonal_estimation_on_a_collection_with_no_seasonal_item_does_nothing():
    result = sn.counter_seasonal_estimate(_panel(seasonal_months=tuple(range(1, 13))))
    assert result.estimated == 0
    assert result.items == ()
    assert any("nothing was estimated" in n for n in result.notes)


# ---------------------------------------------------------------------
# Seasonal adjustment: which engine ran
# ---------------------------------------------------------------------
def test_x13_availability_is_decided_by_looking_for_the_binary(monkeypatch):
    """Not by importing statsmodels' wrapper, which imports perfectly well
    on a machine with no X-13 anywhere. A check that passed on import would
    make the fallback silent, which is the exact failure this guards."""
    monkeypatch.delenv("X13PATH", raising=False)
    monkeypatch.delenv("X12PATH", raising=False)
    monkeypatch.setattr(sn.shutil, "which", lambda name: None)
    available, path, explanation = sn.x13_available()
    assert not available and path is None
    assert "not installed on this machine" in explanation

    monkeypatch.setattr(sn.shutil, "which",
                        lambda name: "/opt/x13/x13as" if name == "x13as" else None)
    available, path, explanation = sn.x13_available()
    assert available and path == "/opt/x13"
    assert "found on PATH" in explanation


def test_x13_path_pointing_at_no_executable_is_reported_as_unavailable(monkeypatch, tmp_path):
    monkeypatch.setenv("X13PATH", str(tmp_path))
    available, _, explanation = sn.x13_available()
    assert not available
    assert "holds no x13as or x12a executable" in explanation


def test_stl_runs_when_x13_is_absent_and_every_output_says_so(monkeypatch):
    """The rule, tested at its source: an STL result presented as "the
    seasonal adjustment" without qualification is a misrepresentation, so
    the label names STL *and* says it is not X-13."""
    monkeypatch.setattr(sn, "x13_available", lambda: (False, None, "no X-13 on this machine"))
    series = _index_series()
    adjustment = sn.adjust(series, SeasonalConfig(adjustment_engine="auto"))

    assert adjustment.engine == "stl"
    assert adjustment.engine_requested == "auto"
    assert adjustment.fell_back
    assert adjustment.fallback_reason == "no X-13 on this machine"
    assert "STL" in adjustment.label
    assert "not X-13ARIMA-SEATS" in adjustment.label
    assert "no X-13 on this machine" in adjustment.label


def test_x13_is_used_and_named_when_the_binary_is_there(monkeypatch):
    """The other branch. This machine has no X-13 binary, so the runner is
    substituted -- which is exactly why `_run_x13` is a module-level
    function with a narrow contract: a fallback whose alternative has never
    been executed is an assumption, not a branch."""
    series = _index_series()
    calls: list[str | None] = []

    def fake_x13(endog, path, periods_per_year):
        calls.append(path)
        adjusted, factors, trend = sn._stl_decompose(endog, periods_per_year, 7)
        return adjusted, factors, trend

    monkeypatch.setattr(sn, "x13_available", lambda: (True, "/opt/x13", "found"))
    monkeypatch.setattr(sn, "_run_x13", fake_x13)
    adjustment = sn.adjust(series, SeasonalConfig(adjustment_engine="auto"))

    assert calls == ["/opt/x13"]
    assert adjustment.engine == "x13"
    assert not adjustment.fell_back
    assert adjustment.label == "seasonally adjusted with X-13ARIMA-SEATS"
    assert "STL" not in adjustment.label


def test_demanding_x13_fails_rather_than_quietly_substituting_stl(monkeypatch):
    monkeypatch.setattr(sn, "x13_available", lambda: (False, None, "no binary here"))
    with pytest.raises(sn.SeasonalError, match="will be labelled as STL"):
        sn.adjust(_index_series(), SeasonalConfig(adjustment_engine="x13"))


def test_x13_failing_mid_run_falls_back_under_auto_and_records_why(monkeypatch):
    def exploding(endog, path, periods_per_year):
        raise RuntimeError("spec file rejected")

    monkeypatch.setattr(sn, "x13_available", lambda: (True, "/opt/x13", "found"))
    monkeypatch.setattr(sn, "_run_x13", exploding)
    adjustment = sn.adjust(_index_series(), SeasonalConfig(adjustment_engine="auto"))
    assert adjustment.engine == "stl"
    assert "spec file rejected" in (adjustment.fallback_reason or "")
    assert "spec file rejected" in adjustment.label

    monkeypatch.setattr(sn, "_run_x13", exploding)
    with pytest.raises(sn.SeasonalError, match="failed on this series"):
        sn.adjust(_index_series(), SeasonalConfig(adjustment_engine="x13"))


def test_asking_for_stl_gets_stl_with_nothing_to_qualify(monkeypatch):
    monkeypatch.setattr(sn, "x13_available", lambda: (True, "/opt/x13", "found"))
    adjustment = sn.adjust(_index_series(), SeasonalConfig(adjustment_engine="stl"))
    assert adjustment.engine == "stl"
    assert not adjustment.fell_back
    assert adjustment.label == "seasonally adjusted with STL (seasonal-trend decomposition " \
                               "by loess)"


# ---------------------------------------------------------------------
# Seasonal adjustment: does it do the right thing
# ---------------------------------------------------------------------
def _index_series(n_years: int = 6, amplitude: float = 0.05, trend: float = 0.003,
                  noise: float = 0.002, seed: int = 4) -> pd.Series:
    """A monthly index with a known seasonal wave and a known trend."""
    rng = np.random.default_rng(seed)
    periods = pd.date_range("2018-01-01", periods=12 * n_years, freq="MS")
    t = np.arange(len(periods))
    wave = amplitude * np.sin(2 * np.pi * (periods.month - 1) / 12)
    values = 100.0 * np.exp(trend * t + wave + rng.normal(0, noise, len(periods)))
    return pd.Series(values, index=periods, name="All items")


def test_a_known_seasonal_pattern_is_removed_without_introducing_a_trend():
    """The acceptance criterion. Seasonal factors that average out over a
    year cannot change a series' underlying growth, so the adjusted and
    unadjusted trends must agree -- and the gap between them is reported in
    percentage points a year whether or not it is small."""
    series = _index_series(n_years=8, amplitude=0.06, trend=0.004)
    adjustment = sn.adjust(series, SeasonalConfig(adjustment_engine="stl"))

    assert adjustment.stability.trend_difference_pp == pytest.approx(0.0, abs=0.1)
    # The wave is gone. Measured on the series detrended by its own centred
    # twelve-month average, because any raw measure of "how much this varies
    # within a year" is dominated by the trend -- a series growing 4.9% a
    # year spans 4.9% from January to December with no seasonality in it at
    # all, and a test that did not remove the trend first would pass on an
    # adjustment that had done nothing.
    def residual_seasonality(values: pd.Series) -> float:
        detrended = values / values.rolling(12, center=True).mean()
        by_month = detrended.dropna().groupby(lambda ts: ts.month).mean()
        return float(by_month.max() - by_month.min())

    assert residual_seasonality(series) > 0.09          # the injected wave, peak to trough
    assert residual_seasonality(adjustment.adjusted) < 0.01
    # and what is left is the injected noise, not a pattern
    assert adjustment.diagnostics["residual_sd_pct"] < 0.5
    # and the trend itself survives
    assert adjustment.stability.trend_adjusted_pct_per_year == pytest.approx(
        (1.004 ** 12 - 1) * 100, rel=0.25)


def test_the_stability_test_re_estimates_the_factors_across_sub_samples():
    adjustment = sn.adjust(_index_series(n_years=9), SeasonalConfig(adjustment_engine="stl"))
    report = adjustment.stability
    assert report.sub_samples >= 2
    assert len(report.factor_ranges) > 0
    assert report.stable
    assert report.max_factor_range_pct <= report.threshold_pct


def test_an_unstable_seasonal_pattern_is_flagged_rather_than_smoothed_over():
    """A factor that moves a great deal between halves of the sample is not
    a seasonal pattern; it is a curve fitted to whatever each stretch did,
    and an adjusted series built on it will be revised."""
    periods = pd.date_range("2016-01-01", periods=12 * 9, freq="MS")
    t = np.arange(len(periods))
    # The wave reverses sign half-way through the sample.
    turning = np.where(t < len(t) / 2, 1.0, -1.0)
    values = 100 * np.exp(0.002 * t + 0.08 * turning * np.sin(2 * np.pi * (periods.month - 1) / 12))
    adjustment = sn.adjust(pd.Series(values, index=periods),
                           SeasonalConfig(adjustment_engine="stl", stability_threshold_pct=2.0))
    assert not adjustment.stability.stable
    assert adjustment.stability.max_factor_range_pct > 2.0
    assert any("not stable across this sample" in w for w in adjustment.warnings)


def test_a_series_shorter_than_two_cycles_is_refused():
    short = _index_series(n_years=1)
    with pytest.raises(sn.SeasonalError, match="at least two full years"):
        sn.adjust(short, SeasonalConfig(adjustment_engine="stl"))


def test_the_adjustment_carries_the_unadjusted_series_on_the_same_object():
    """The mechanism behind "published alongside, everywhere": there is no
    code path that can obtain the adjusted series without also holding what
    it was adjusted from."""
    adjustment = sn.adjust(_index_series(), SeasonalConfig(adjustment_engine="stl"))
    assert list(adjustment.frame.columns[:3]) == ["unadjusted", "adjusted", "seasonal_factor"]
    assert adjustment.unadjusted.notna().all()
    assert len(adjustment.unadjusted) == len(adjustment.adjusted)
    assert not adjustment.unadjusted.equals(adjustment.adjusted)


# ---------------------------------------------------------------------
# The stage, and the method note
# ---------------------------------------------------------------------
def test_the_seasonal_stage_runs_everything_it_is_asked_for(compiled):
    cfg = SeasonalConfig(enabled=True, rothwell=True, counter_seasonal=True)
    result = sn.run_seasonal(compiled["imputed"], compiled["indices"], cfg,
                             compiled["config"].index)
    assert result.seasonal_item_count >= 3
    assert result.comparison is not None and len(result.comparison.results) == 2
    assert result.adjustment is not None
    assert result.rothwell is not None
    assert result.counter_seasonal is not None
    assert "X-13ARIMA-SEATS" in result.engine_label


def test_a_stage_that_could_not_run_says_why_rather_than_producing_nothing(compiled):
    cfg = SeasonalConfig(enabled=True, adjustment_series="Not a column")
    result = sn.run_seasonal(compiled["imputed"], compiled["indices"], cfg,
                             compiled["config"].index)
    assert result.adjustment is None
    assert any("no column named 'Not a column'" in n for n in result.notes)
    assert result.engine_label == "no seasonal adjustment was computed"


def test_a_single_treatment_run_says_the_other_was_not_computed(compiled):
    cfg = SeasonalConfig(enabled=True, treatment="class_confinement", adjust=False)
    result = sn.run_seasonal(compiled["imputed"], compiled["indices"], cfg,
                             compiled["config"].index)
    assert result.comparison is not None
    assert set(result.comparison.results) == {"class_confinement"}
    assert any("set treatment to 'both' to see it" in n for n in result.comparison.notes)


def test_the_method_note_names_the_engine_that_ran(compiled):
    cfg = SeasonalConfig(enabled=True)
    result = sn.run_seasonal(compiled["imputed"], compiled["indices"], cfg,
                             compiled["config"].index)
    note = sn.adjustment_note(result)
    assert "X-13ARIMA-SEATS" in note
    assert "unadjusted series is published alongside it" in note
    assert "strictly seasonal" in note
    assert sn.adjustment_note(None) == ""


def test_the_config_rejects_an_engine_or_treatment_it_does_not_have():
    for kwargs, match in (
            ({"adjustment_engine": "seats"}, "adjustment_engine"),
            ({"treatment": "ignore"}, "treatment"),
            ({"periods_per_year": 1}, "no structure"),
            ({"min_years": 1}, "at least two years"),
            ({"stl_seasonal": 6}, "odd number"),
            ({"stability_sub_samples": 1}, "at least two sub-samples")):
        with pytest.raises(ValueError, match=match):
            SeasonalConfig(**kwargs)


# ---------------------------------------------------------------------
# The unadjusted series travels with the adjusted one, into every format
# ---------------------------------------------------------------------
@pytest.fixture(scope="module")
def adjusted_run(collection: pd.DataFrame) -> dict:
    return run_pipeline(collection, RunConfig(
        seasonal=SeasonalConfig(enabled=True, rothwell=True), label="seasonal run"))


def test_the_publication_table_carries_both_series_with_the_engine_named(adjusted_run):
    from pricelab.reporting import exports

    table = exports.publication_table(adjusted_run)
    names = set(table["series"])
    assert "seasonally adjusted: All items" in names
    assert "unadjusted: All items" in names

    basis = dict(zip(table["series"], table["basis"], strict=True))
    assert "X-13ARIMA-SEATS" in basis["seasonally adjusted: All items"]
    assert basis["unadjusted: All items"] == "as compiled, before seasonal adjustment"
    # the ordinary bilateral series carry no basis, because theirs is the
    # run's own configuration and is in the stamp already
    assert basis["All items"] == ""


def test_the_stamped_csv_carries_both_series(adjusted_run):
    from pricelab.core.provenance import build_stamp
    from pricelab.reporting import exports

    csv = exports.index_csv(adjusted_run, build_stamp(adjusted_run, "seasonal run"))
    assert "seasonally adjusted: All items" in csv
    assert "unadjusted: All items" in csv
    assert "X-13ARIMA-SEATS" in csv


def test_the_markdown_report_names_the_engine_and_shows_both_series(adjusted_run):
    markdown = build_markdown(adjusted_run, build_narrative(adjusted_run), "seasonal run")
    assert "## Seasonal adjustment" in markdown
    assert "## Strictly seasonal item treatment" in markdown
    assert "## Rothwell index" in markdown
    assert "X-13ARIMA-SEATS" in markdown
    assert "unadjusted" in markdown and "adjusted" in markdown


def test_the_word_report_and_the_excel_pack_carry_the_same_sections(adjusted_run, collection):
    from docx import Document
    from openpyxl import load_workbook

    from pricelab import build_all_charts, build_docx
    from pricelab.core.provenance import build_stamp
    from pricelab.reporting.excel import build_evidence_pack

    stamp = build_stamp(adjusted_run, "seasonal run")
    nar = build_narrative(adjusted_run)
    document = Document(io.BytesIO(
        build_docx(adjusted_run, nar, build_all_charts(adjusted_run), "seasonal run", stamp)))
    headings = [p.text for p in document.paragraphs if p.style.name == "Heading 2"]
    assert "Seasonal adjustment" in headings
    assert any("X-13ARIMA-SEATS" in p.text for p in document.paragraphs)

    book = load_workbook(io.BytesIO(
        build_evidence_pack(adjusted_run, stamp, source=collection)))
    assert "Seasonal adjustment" in book.sheetnames
    assert "Series basis" in book.sheetnames
    basis = [[c.value for c in row] for row in book["Series basis"].iter_rows()]
    assert any(row[0] and "seasonally adjusted" in str(row[0]) for row in basis)
    assert any(row[1] and "X-13ARIMA-SEATS" in str(row[1]) for row in basis)


def test_the_deck_carries_the_engine_too(adjusted_run):
    from pptx import Presentation

    from pricelab import build_all_charts, build_deck

    deck = build_deck(adjusted_run, build_narrative(adjusted_run),
                      build_all_charts(adjusted_run), "seasonal run")
    slides = Presentation(io.BytesIO(deck)).slides
    text = "\n".join("\n".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)
                     for slide in slides)
    assert "What produced these series" in text
    assert "X-13ARIMA-SEATS" in text


# ---------------------------------------------------------------------
# The page
# ---------------------------------------------------------------------
class FakeUpload:
    def __init__(self, name: str, data: bytes) -> None:
        self.name, self._data, self.size = name, data, len(data)

    def getvalue(self) -> bytes:
        return self._data


@pytest.fixture()
def deployment(tmp_path, monkeypatch):
    monkeypatch.setenv("PRICELAB_DATABASE_URL", database_url(tmp_path, "seasonal.db"))
    monkeypatch.setenv("PRICELAB_STORE_DIR", str(tmp_path / "store"))
    get_settings.cache_clear()
    db.reset_db_state()
    from pricelab.core import ledger, registry, security  # noqa: F401  register tables
    from pricelab.data import classification, mapping, validation  # noqa: F401
    db.init_db()
    from pricelab.core.cache import reset_analysis_cache
    from pricelab.core.ratelimit import reset_upload_limiter
    reset_analysis_cache()
    reset_upload_limiter()
    yield tmp_path
    db.reset_db_state()
    get_settings.cache_clear()


def _ingest(monkeypatch) -> dict:
    monkeypatch.setattr(streamlit, "file_uploader",
                        lambda *a, **k: FakeUpload("prices.xlsx", PRICES.read_bytes()))
    script = """
import streamlit as st
from pricelab.core.models import Role
from pricelab.core.security import set_current_role
import pages.ingest as page
set_current_role(Role.COMPILER)
st.session_state.setdefault("pricelab_username", "tester1")
page.render()
"""
    at = AppTest.from_string(script, default_timeout=300)
    at.run()
    assert not at.exception, at.exception
    next(b for b in at.button if b.label == "Confirm column mapping").click().run()
    assert not at.exception, at.exception
    monkeypatch.setattr(streamlit, "file_uploader", lambda *a, **k: None)
    return {k: v for k, v in at.session_state.filtered_state.items()
            if not str(k).startswith("$$")}


def _page(state: dict, role: str = "COMPILER") -> AppTest:
    script = f"""
import streamlit as st
from pricelab.core.models import Role
from pricelab.core.security import set_current_role
import pages.seasonal as page
set_current_role(Role.{role})
st.session_state.setdefault("pricelab_username", "tester1")
page.render()
"""
    at = AppTest.from_string(script, default_timeout=300)
    for key, value in state.items():
        at.session_state[key] = value
    at.run()
    assert not at.exception, at.exception
    return at


def _text(at: AppTest) -> str:
    return "\n".join(m.value for m in at.markdown) + "\n" + "\n".join(c.value for c in at.caption)


def test_the_page_finds_the_seasonal_items_and_reports_the_treatment_gap(
        deployment, monkeypatch):
    at = _page(_ingest(monkeypatch))
    text = _text(at)
    assert "Strictly seasonal items" in text
    assert "Class confinement against weight update" in text

    labels = [m.label for m in at.metric]
    assert "Items off the shelf for part of every year" in labels
    assert "Largest gap" in labels
    gap = next(m for m in at.metric if m.label == "Largest gap")
    assert float(str(gap.value).split()[0]) > 0.01


def test_the_page_adjusts_and_names_the_engine_on_screen(deployment, monkeypatch):
    """The page-level half of the rule: a reader who never opens a download
    has still been told what produced the line, before the chart."""
    at = _page(_ingest(monkeypatch))
    next(b for b in at.button if b.label == "Adjust").click().run()
    assert not at.exception, at.exception

    adjustment = at.session_state["sn_adjustment"]
    assert adjustment.engine == "stl"
    assert adjustment.fell_back

    surfaced = ([w.value for w in at.warning] + [s.value for s in at.success]
                + [c.value for c in at.caption])
    assert any("X-13ARIMA-SEATS" in t for t in surfaced)
    assert any("not X-13ARIMA-SEATS" in t for t in surfaced)
    assert any("Unadjusted and adjusted on the same axes" in c.value for c in at.caption)
    assert "Trend introduced by the adjustment" in [m.label for m in at.metric]


def test_the_page_computes_the_rothwell_index_and_counter_seasonal_estimates(
        deployment, monkeypatch):
    state = _ingest(monkeypatch)
    at = _page(state)
    next(b for b in at.button if b.label == "Compute the Rothwell index").click().run()
    assert not at.exception, at.exception
    assert at.session_state["sn_rothwell_result"].items_in_base > 0

    next(b for b in at.button if b.label == "Estimate off-season prices").click().run()
    assert not at.exception, at.exception
    estimate = at.session_state["sn_counter_result"]
    assert estimate.estimated > 0
    assert "Off-season prices estimated" in [m.label for m in at.metric]


def test_the_page_computes_the_seasonal_multilateral_forms(deployment, monkeypatch):
    """Phase 5's orphan, folded in here: the year-over-year and rolling-year
    forms are reachable from the product rather than living in the library
    alone."""
    at = _page(_ingest(monkeypatch))
    at.session_state["sn_ml_method"] = "tpd"       # prices only; no quantities in this file
    next(b for b in at.button if b.label == "Compute the seasonal forms").click().run()
    assert not at.exception, at.exception
    result = at.session_state["sn_seasonal_multilateral"]
    assert result.months_compiled == 12
    assert result.year_over_year.notna().any()
    assert result.rolling_year.notna().any()


def test_a_viewer_is_refused_the_seasonality_page(deployment):
    import pages.seasonal as page
    from pricelab.core.models import Role
    from pricelab.core.security import AccessDenied, set_current_role

    set_current_role(Role.VIEWER)
    try:
        with pytest.raises(AccessDenied, match="administrator"):
            page.render()
    finally:
        set_current_role(Role.COMPILER)
