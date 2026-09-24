"""The slide deck reads as finished work: no slide whose only content is why
it is empty, no prose that starts mid-sentence, no field cut off, and every
finding in the same two-column layout."""

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from pathlib import Path

import pandas as pd
import pytest
from pptx import Presentation
from pptx.util import Emu
from test_release_contributions import _collection

from pricelab import analyse, infer_schema, run_pipeline, standardise
from pricelab.core.config import RunConfig
from pricelab.core.provenance import build_stamp
from pricelab.engine import decomposition as dc
from pricelab.reporting.deck import build_deck, provenance_display_value
from pricelab.reporting.report import CONTRIBUTIONS_TITLE

FIXTURE = Path(__file__).resolve().parents[1] / "supermarket_price_collection.xlsx"
EYEBROWS = {"DATA QUALITY", "SAMPLE STRUCTURE", "PRICE TREND", "SEASONALITY", "METHOD"}


@dataclass
class Frame:
    text: str
    left: float
    top: float
    width: float
    height: float


@dataclass
class Slide:
    frames: list[Frame]
    pictures: int
    tables: int

    @property
    def title(self) -> str:
        return self.frames[0].text if self.frames else ""


def _inches(value: int) -> float:
    return float(Emu(value).inches)


def _slides(data: bytes) -> list[Slide]:
    slides = []
    for slide in Presentation(io.BytesIO(data)).slides:
        frames = [Frame(sh.text_frame.text.strip(), _inches(sh.left), _inches(sh.top),
                        _inches(sh.width), _inches(sh.height))
                  for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        slides.append(Slide(frames, sum(1 for sh in slide.shapes if sh.shape_type == 13),
                            sum(1 for sh in slide.shapes if sh.has_table)))
    return slides


def _deck(res: dict, label: str) -> tuple[bytes, dict]:
    from pricelab import build_narrative
    from pricelab.reporting.charts import build_all_charts

    stamp = build_stamp(res, label)
    return build_deck(res, build_narrative(res), build_all_charts(res), label, stamp), \
        {"stamp": stamp}


@pytest.fixture(scope="module")
def unweighted() -> tuple[bytes, dict]:
    if not FIXTURE.exists():
        pytest.skip("fixture workbook not present")
    raw = pd.read_excel(FIXTURE)
    out = analyse(standardise(raw, infer_schema(raw)), "fixture", RunConfig())
    return _deck(out["result"], "fixture")


@pytest.fixture(scope="module")
def weighted() -> tuple[bytes, dict]:
    return _deck(run_pipeline(_collection(weighted=True), RunConfig(label="weighted")),
                 "weighted")


DECKS = ["unweighted", "weighted"]
#: Prose opens with an ordinary word: letters only, a trailing comma or the
#: like allowed. "python-3.14.0;" or a commit hash is a data value.
_PROSE_START = re.compile(r"^[a-z][a-z']{0,19}[,.;:]?$")


@pytest.mark.parametrize("deck", DECKS)
def test_no_prose_on_any_slide_begins_with_a_lowercase_letter(deck, request):
    """Prose -- four words or more, opening with an ordinary word -- starts
    with a capital. Data values (a run label, a category, a commit hash) are
    not prose and are exempt."""
    data, _ = request.getfixturevalue(deck)
    offenders = [(i, f.text[:80]) for i, s in enumerate(_slides(data), 1) for f in s.frames
                 if len(f.text.split()) >= 4 and _PROSE_START.match(f.text.split()[0])]
    assert not offenders, offenders


@pytest.mark.parametrize("deck", DECKS)
def test_no_titled_slide_is_only_a_reason_string(deck, request):
    """A deck never contains a slide whose only content, beyond its title, is
    one block of text explaining why it has nothing to show."""
    data, _ = request.getfixturevalue(deck)
    thin = [(i, s.title) for i, s in enumerate(_slides(data), 1)
            if len(s.frames) <= 2 and not s.pictures and not s.tables]
    assert not thin, thin


def test_an_unweighted_run_drops_the_contributions_slide_and_says_why_on_the_method_slide(
        unweighted):
    data, _ = unweighted
    slides = _slides(data)
    assert CONTRIBUTIONS_TITLE not in [s.title for s in slides]
    method = next(s for s in slides if s.title == "How this was produced")
    limitations = " ".join(f.text for f in method.frames)
    assert "This run has no expenditure weights" in limitations
    assert "no exact additive decomposition" in limitations
    # and the headline slide no longer contradicts a weighted run (below)
    stats = next(s for s in slides if s.title == "The headline numbers")
    assert any("equally weighted" in f.text for f in stats.frames)


def test_a_weighted_run_keeps_the_contributions_slide_and_says_it_is_weighted(weighted):
    data, _ = weighted
    slides = _slides(data)
    assert CONTRIBUTIONS_TITLE in [s.title for s in slides]
    method = next(s for s in slides if s.title == "How this was produced")
    assert any("Expenditure weights supplied" in f.text for f in method.frames)
    assert not any("No expenditure weights" in f.text for s in slides for f in s.frames)
    stats = next(s for s in slides if s.title == "The headline numbers")
    assert any("weighted by the expenditure weights supplied" in f.text for f in stats.frames)


def test_the_reason_contributions_are_absent_is_a_complete_sentence():
    res = run_pipeline(_collection(weighted=False), RunConfig())
    (note,) = dc.components_from_run(res).notes
    assert note[0].isupper() and note.endswith(".")
    assert "rather than computed for a different aggregate" not in note


@pytest.mark.parametrize("deck", DECKS)
def test_the_provenance_slide_carries_every_field_in_full_and_on_the_slide(deck, request):
    """Previously cut at 160 characters, mid-word ("whe"). Every value is
    shown whole, and the last row still ends above the slide's edge."""
    data, extra = request.getfixturevalue(deck)
    prov = next(s for s in _slides(data) if s.title == "Provenance")
    shown = {f.text for f in prov.frames}
    for field, value in extra["stamp"].rows():
        if field == "Parameters (JSON)":
            continue
        assert provenance_display_value(field, value) in shown, field
    assert max(f.top + f.height for f in prov.frames) <= 7.5


@pytest.mark.parametrize("deck", DECKS)
def test_every_finding_uses_the_two_column_layout(deck, request):
    """A finding with a chart has it on the right; one without has its
    evidence and recommended action there instead, never a full-width block
    of text."""
    data, _ = request.getfixturevalue(deck)
    findings = [s for s in _slides(data) if s.frames and s.frames[0].text in EYEBROWS]
    assert findings
    for slide in findings:
        assert slide.pictures or any(f.text == "THE EVIDENCE" for f in slide.frames), \
            slide.frames[1].text
        left_prose = [f for f in slide.frames if f.left < 1.0]
        assert all(f.width <= 5.0 for f in left_prose), [f.width for f in left_prose]
