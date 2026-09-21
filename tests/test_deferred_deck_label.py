"""The "= 100" label must name the period the series was actually rebased to.

Held open as a `strict=True` xfail from the Phase 1 patch until Phase 3,
which is the phase that made the bug reachable: Lowe and Young need a real,
independently-set reference period, and the deck is the artefact that leaves
this application to be read by someone who cannot check it against the
source data.

Fixed by giving the rebasing rule one home --
`engine.index.resolve_index_reference_period` -- and having the arithmetic
(`build_index`'s final rescaling) and every label that describes it (this
chart's y-axis, the deck's headline stat, the report's method note, the
Findings metric, the Index build reference-period panel) read that one
function. A label can no longer name a different period from the one the
series was rebased to, because neither side derives it independently any
more. The xfail marker is gone; this is now an ordinary regression test.
"""

from io import BytesIO

import pandas as pd
from pptx import Presentation

from pricelab import run_pipeline
from pricelab.core.config import IndexConfig, RunConfig
from pricelab.engine.insights import build_narrative
from pricelab.reporting.charts import build_all_charts
from pricelab.reporting.deck import build_deck


def _panel_with_a_late_index_reference_period():
    periods = pd.date_range("2020-01-01", periods=4, freq="MS")  # Jan-Apr 2020
    df = pd.DataFrame({
        "period": list(periods) * 2,
        "category": ["Bread"] * 4 + ["Milk"] * 4,
        "item_id": ["1"] * 4 + ["2"] * 4,
        "item_name": ["White"] * 4 + ["Semi"] * 4,
        "price_reported": [1.0, 1.05, 1.10, 1.15, 2.0, 1.9, 2.2, 2.6],
    })
    # Rebased to March 2020 (periods[2]), not the series' first period.
    index_ref = periods[2]
    cfg = RunConfig(index=IndexConfig(chained=True, index_reference_period=str(index_ref.date())))
    return df, cfg, index_ref, periods[0]


def test_chart_ylabel_and_deck_text_report_the_actual_index_reference_period():
    df, cfg, index_ref, first_period = _panel_with_a_late_index_reference_period()
    res = run_pipeline(df, cfg)
    assert "indices" in res, "fixture must produce a valid index for this test to mean anything"

    nar = build_narrative(res)
    charts = build_all_charts(res)

    correct_label_fragment = f"{index_ref:%b %Y} = 100"
    wrong_label_fragment = f"{first_period:%b %Y} = 100"
    assert correct_label_fragment != wrong_label_fragment, "test data must make the two differ"

    # reporting/charts.py's index_chart
    chart_ylabel = charts["index"].axes[0].get_ylabel()
    assert correct_label_fragment in chart_ylabel, (
        f"chart y-axis label {chart_ylabel!r} does not name the configured "
        f"index_reference_period ({index_ref:%b %Y})")

    # reporting/deck.py's headline stat callout
    deck_bytes = build_deck(res, nar, charts, label="deferred label test")
    prs = Presentation(BytesIO(deck_bytes))
    all_text = "\n".join(
        shape.text_frame.text
        for slide in prs.slides
        for shape in slide.shapes
        if shape.has_text_frame
    )
    assert correct_label_fragment in all_text, (
        "no slide names the configured index_reference_period "
        f"({index_ref:%b %Y}) in a '= 100' label")


def test_the_method_note_names_the_actual_index_reference_period():
    """The written report's method note carried the same hardcode as the
    chart and the deck, and is the artefact a reviewer reads to check the
    method, so it gets the same guarantee."""
    from pricelab.reporting.report import method_note

    df, cfg, index_ref, first_period = _panel_with_a_late_index_reference_period()
    res = run_pipeline(df, cfg)

    note = method_note(res)
    assert f"{index_ref:%B %Y}" in note
    assert f"set to 100 at {first_period:%B %Y}" not in note
