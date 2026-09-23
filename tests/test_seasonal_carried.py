"""Phase 7a, Task 0: the three items carried over from seasonal adjustment.

a. Every surface an adjusted series reaches is enumerated in code
   (`reporting.exports.SEASONAL_SURFACES`), and this module renders each one
   and checks that it names the engine that ran, shows the unadjusted series
   beside the adjusted one, and states that the adjustment was direct and
   what that means for the parts. A module that touches the adjusted series
   without being enumerated fails the inventory test below.
b. A series with no seasonality, adjusted, comes back essentially unchanged,
   to a tolerance stated and justified in the test.
c. The additivity gap is labelled, not closed: a directly adjusted total and
   separately adjusted components need not agree, and every surface says so.
"""

from __future__ import annotations

import io
import re
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
import pytest
from dbtarget import database_url
from streamlit.testing.v1 import AppTest

from pricelab import (
    build_all_charts,
    build_deck,
    build_docx,
    build_markdown,
    build_narrative,
    infer_schema,
    run_pipeline,
    standardise,
)
from pricelab.core import db
from pricelab.core.config import RunConfig, SeasonalConfig, get_settings
from pricelab.core.provenance import build_stamp
from pricelab.engine import seasonal as sn
from pricelab.reporting import exports
from pricelab.reporting.charts import seasonal_adjustment_chart
from pricelab.reporting.excel import build_evidence_pack

REPO_ROOT = Path(__file__).resolve().parents[1]
PRICES = REPO_ROOT / "supermarket_price_collection.xlsx"
needs_prices = pytest.mark.skipif(not PRICES.exists(), reason="fixture workbook not present")

#: What every surface must carry, as phrases that survive being wrapped,
#: hyphenated or split across a PDF's text lines once whitespace is collapsed.
ENGINE = "STL"                       # the engine that runs here (no X-13 binary)
FALLBACK = "not X-13ARIMA-SEATS"     # and the statement that it is the fallback
UNADJUSTED = "unadjusted"
DIRECT = "adjusted directly"
ADDITIVITY = "need not sum"


def _normalise(text: str) -> str:
    return re.sub(r"\s+", " ", text.replace("-\n", "")).strip()


def _assert_discloses(surface: str, text: str) -> None:
    flat = _normalise(text)
    lowered = flat.lower()
    missing = [phrase for phrase, present in (
        (ENGINE, ENGINE in flat), (FALLBACK, FALLBACK in flat),
        (UNADJUSTED, UNADJUSTED in lowered), (DIRECT, DIRECT in lowered),
        (ADDITIVITY, ADDITIVITY in lowered)) if not present]
    assert not missing, f"the {surface} surface omits {missing}"


# ---------------------------------------------------------------------
# a. Every surface, rendered and read back
# ---------------------------------------------------------------------
@pytest.fixture(scope="module")
def collection() -> pd.DataFrame:
    if not PRICES.exists():
        pytest.skip("fixture workbook not present")
    raw = pd.read_excel(PRICES)
    return standardise(raw, infer_schema(raw))


@pytest.fixture(scope="module")
def adjusted(collection: pd.DataFrame) -> dict[str, Any]:
    res = run_pipeline(collection, RunConfig(
        seasonal=SeasonalConfig(enabled=True, adjustment_engine="auto"), label="seasonal run"))
    assert res["seasonal"].adjustment is not None
    assert res["seasonal"].adjustment.engine == "stl", "these tests assume no X-13 binary"
    return res


def _figure_text(fig: Any) -> str:
    parts = [t.get_text() for t in fig.texts]
    for ax in fig.axes:
        parts += [ax.get_title(loc="left"), ax.get_title(), ax.get_ylabel()]
        legend = ax.get_legend()
        if legend is not None:
            parts += [t.get_text() for t in legend.get_texts()]
    return "\n".join(parts)


def _pptx_text(data: bytes) -> str:
    from pptx import Presentation

    return "\n".join("\n".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)
                     for slide in Presentation(io.BytesIO(data)).slides)


def _docx_text(data: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(data))
    cells = [c.text for t in document.tables for row in t.rows for c in row.cells]
    return "\n".join([p.text for p in document.paragraphs] + cells)


def _xlsx_text(data: bytes) -> str:
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(data), read_only=True)
    return "\n".join(str(v) for name in ("Seasonal adjustment", "Series basis")
                     for row in book[name].iter_rows(values_only=True) for v in row
                     if v is not None)


def _bulletin_text(res: dict[str, Any], tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
                   collection: pd.DataFrame) -> str:
    from pypdf import PdfReader

    from pricelab.core.registry import register_run
    from pricelab.reporting import bulletin

    monkeypatch.setenv("PRICELAB_DATABASE_URL", database_url(tmp_path, "surfaces.db"))
    monkeypatch.setenv("PRICELAB_STORE_DIR", str(tmp_path / "store"))
    get_settings.cache_clear()
    db.reset_db_state()
    db.init_db()
    try:
        with db.session_scope() as s:
            run = register_run(s, collection, res["config"], "seasonal run", result=res)
            stamp = build_stamp(res, "seasonal run", run=run)
            pdf = bulletin.build_bulletin(res, build_narrative(res), build_all_charts(res),
                                          stamp, run=run)
    finally:
        db.reset_db_state()
        get_settings.cache_clear()
    return "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(pdf)).pages)


def _page_text(res: dict[str, Any]) -> str:
    script = """
import streamlit as st
from pricelab.core.models import Role
from pricelab.core.security import set_current_role
import pages.seasonal as page
set_current_role(Role.COMPILER)
st.session_state.setdefault("pricelab_username", "tester1")
page.render()
"""
    at = AppTest.from_string(script, default_timeout=300)
    at.session_state["analysis"] = {"result": res, "narrative": None, "decisions": [],
                                    "label": "seasonal run"}
    at.session_state["sn_engine"] = "auto"
    at.run()
    assert not at.exception, at.exception
    next(b for b in at.button if b.label == "Adjust").click().run()
    assert not at.exception, at.exception
    shown = ([w.value for w in at.warning] + [s.value for s in at.success]
             + [c.value for c in at.caption])
    # The page's own chart is the shared figure; it is read back directly
    # below under "chart", so here the page must at least say, before and
    # under the chart, what produced it and that both series are drawn.
    return "\n".join(shown)


def test_every_enumerated_surface_names_the_engine_and_carries_both_series(
        adjusted, collection, tmp_path, monkeypatch):
    res = adjusted
    stamp = build_stamp(res, "seasonal run")
    nar = build_narrative(res)
    charts = build_all_charts(res)
    rendered: dict[str, str] = {
        "page": _page_text(res),
        "chart": _figure_text(charts["seasonal_adjustment"]),
        "method_note": sn.adjustment_note(res["seasonal"]),
        "markdown": build_markdown(res, nar, "seasonal run"),
        "word": _docx_text(build_docx(res, nar, charts, "seasonal run", stamp)),
        "deck": _pptx_text(build_deck(res, nar, charts, "seasonal run")),
        "excel": _xlsx_text(build_evidence_pack(res, stamp, source=collection)),
        "bulletin": _bulletin_text(res, tmp_path, monkeypatch, collection),
        "csv": exports.index_csv(res, stamp),
        "sdmx": exports.to_sdmx_ml(res, stamp).decode("utf-8"),
    }
    # Every enumerated surface was rendered, and nothing was rendered that
    # the enumeration does not know about.
    assert set(rendered) == set(exports.SEASONAL_SURFACES)
    for surface, text in rendered.items():
        _assert_discloses(surface, text)


def test_the_chart_draws_both_series_not_only_the_adjusted_one(adjusted):
    fig = seasonal_adjustment_chart(adjusted["seasonal"].adjustment)
    labels = [str(line.get_label()) for line in fig.axes[0].lines]
    assert any(label.startswith("Unadjusted") for label in labels)
    assert any(label.startswith("Seasonally adjusted (STL") for label in labels)


def test_the_sdmx_basis_is_on_the_adjusted_series_itself(adjusted):
    from lxml import etree

    stamp = build_stamp(adjusted, "seasonal run")
    root = etree.fromstring(exports.to_sdmx_ml(adjusted, stamp))
    generic = "{http://www.sdmx.org/resources/sdmxml/schemas/v2_1/data/generic}"
    found: dict[str, str] = {}
    for series in root.iter(f"{generic}Series"):
        key = series.find(f"{generic}SeriesKey/{generic}Value").get("value")
        basis = [v.get("value") for v in series.iter(f"{generic}Value") if v.get("id") == "BASIS"]
        found[key] = basis[0] if basis else ""
    assert FALLBACK in found["seasonally adjusted: All items"]
    assert "adjusted directly" in found["seasonally adjusted: All items"]
    assert found["unadjusted: All items"] == "as compiled, before seasonal adjustment"
    assert found["All items"] == ""        # the ordinary series carry no basis
    ok, errors = exports.validate_sdmx_ml(exports.to_sdmx_ml(adjusted, stamp),
                                          str(REPO_ROOT / "tests" / "fixtures" / "sdmx_2_1"))
    assert ok, errors


#: Which surface(s) each module that handles the adjusted series feeds. The
#: inventory test fails when a module outside this map starts touching the
#: adjusted series: it is a new surface, and it has to be enumerated -- and so
#: checked -- before it can ship.
MODULE_SURFACES = {
    "pages/seasonal.py": {"page"},
    "pricelab/reporting/charts.py": {"chart"},
    "pricelab/reporting/report.py": {"method_note", "markdown", "word"},
    "pricelab/reporting/deck.py": {"deck"},
    "pricelab/reporting/excel.py": {"excel"},
    "pricelab/reporting/bulletin.py": {"bulletin"},
    "pricelab/reporting/exports.py": {"csv", "sdmx"},
}
TOUCHES_ADJUSTED = re.compile(r"\.adjustment\b|seasonal_adjustment\b|SEASONAL_ADJUSTED_PREFIX"
                              r"|sn\.adjust\(|seasonally adjusted: ")


def test_every_module_that_handles_the_adjusted_series_is_an_enumerated_surface():
    touching = sorted(
        str(path.relative_to(REPO_ROOT)).replace("\\", "/")
        for folder in ("pages", "pricelab/reporting")
        for path in (REPO_ROOT / folder).glob("*.py")
        if TOUCHES_ADJUSTED.search(path.read_text(encoding="utf-8")))
    unlisted = [module for module in touching if module not in MODULE_SURFACES]
    assert not unlisted, (
        f"{unlisted} handle the seasonally adjusted series but are not enumerated in "
        "reporting.exports.SEASONAL_SURFACES; add the surface there and a renderer here")
    assert set().union(*MODULE_SURFACES.values()) == set(exports.SEASONAL_SURFACES)


# ---------------------------------------------------------------------
# b. No seasonality in, essentially nothing out
# ---------------------------------------------------------------------
def _no_season(n_years: int, sigma: float, seed: int, trend: float = 0.003) -> pd.Series:
    rng = np.random.default_rng(seed)
    periods = pd.date_range("2015-01-01", periods=12 * n_years, freq="MS")
    t = np.arange(len(periods))
    return pd.Series(100.0 * np.exp(trend * t + rng.normal(0.0, sigma, len(periods))),
                     index=periods, name="All items")


def test_a_trend_with_no_season_and_no_noise_comes_back_unchanged():
    """With nothing but a trend there is nothing to remove, and the adjusted
    series must equal the unadjusted one to floating-point precision. The
    tolerance is 1e-10 in log terms -- about 1e-8 percent -- because anything
    above rounding would be a seasonal pattern invented from a straight line.
    (Measured: 1.7e-14.)"""
    series = _no_season(6, 0.0, 0)
    adjustment = sn.adjust(series, SeasonalConfig(adjustment_engine="stl"))
    change = np.log(adjustment.adjusted / adjustment.unadjusted)
    assert float(change.abs().max()) < 1e-10
    assert abs(adjustment.stability.trend_difference_pp) < 1e-8


@pytest.mark.parametrize("seed", range(10))
def test_a_series_with_no_season_is_moved_by_less_than_its_own_noise(seed):
    """The tolerance, and why it is the right one.

    A noisy series with no seasonality cannot be returned *exactly*
    unchanged: STL estimates each calendar month's factor from a handful of
    that month's observations, so its estimate carries sampling error of
    roughly sigma / sqrt(observations per month), and dividing it out moves
    each point by a fraction of the noise. What it must not do is move the
    series by more than the noise itself -- a "seasonal" pattern larger than
    the irregular, found in data with none, would be manufactured.

    So the yardstick is the irregular's own standard deviation, sigma, not a
    fixed percentage: a fixed 0.1% would be failed by a correct adjustment
    of a noisy series and passed by a broken one on a smooth series. The
    test: the root-mean-square change the adjustment makes, in log terms, is
    at most sigma, and the trend moves by less than 0.1 points a year (the
    threshold at which `adjust` warns). Measured over 40 seeds of ten years:
    median 0.62 sigma, worst 0.77 sigma; the ten seeds here are a sample of
    those, not a chosen few. On only four years the worst of 40 seeds
    reached 1.13 sigma -- short samples fit larger spurious factors, which
    is recorded in docs/methodology/seasonal.md rather than hidden by
    testing only the long case.
    """
    sigma = 0.004
    adjustment = sn.adjust(_no_season(10, sigma, seed), SeasonalConfig(adjustment_engine="stl"))
    change = np.log(adjustment.adjusted / adjustment.unadjusted)
    rms = float(np.sqrt((change ** 2).mean()))
    assert rms <= sigma, f"the adjustment moved a season-free series by {rms / sigma:.2f} sigma"
    assert abs(adjustment.stability.trend_difference_pp) < 0.1


def test_the_null_tolerance_would_catch_a_real_season_being_left_in_or_invented():
    """The tolerance discriminates: a genuine 5% seasonal wave, adjusted,
    changes the series by far more than sigma -- so a test passing at
    'at most sigma' is saying something about the absence of a season, not
    passing everything."""
    sigma = 0.004
    periods = pd.date_range("2015-01-01", periods=120, freq="MS")
    rng = np.random.default_rng(0)
    wave = 0.05 * np.sin(2 * np.pi * (periods.month - 1) / 12)
    series = pd.Series(100 * np.exp(0.003 * np.arange(120) + wave + rng.normal(0, sigma, 120)),
                       index=periods)
    adjustment = sn.adjust(series, SeasonalConfig(adjustment_engine="stl"))
    rms = float(np.sqrt((np.log(adjustment.adjusted / adjustment.unadjusted) ** 2).mean()))
    assert rms > 5 * sigma


# ---------------------------------------------------------------------
# c. The additivity gap, labelled
# ---------------------------------------------------------------------
def test_the_label_says_direct_and_warns_the_parts_need_not_sum():
    adjustment = sn.adjust(_no_season(6, 0.002, 1), SeasonalConfig(adjustment_engine="stl"),
                           series_name="Food")
    assert adjustment.approach == "direct"
    assert "adjusted directly (the Food series itself" in adjustment.label
    assert "need not sum to an adjusted total" in adjustment.label
    assert "no constraint forcing them to has been applied" in adjustment.label


@needs_prices
def test_directly_adjusted_components_do_not_sum_to_the_directly_adjusted_total(collection):
    """Why the label is needed, shown on real data: adjust two categories and
    their equally weighted aggregate directly, and the aggregate of the
    adjusted categories is not the adjusted aggregate. The gap is reported
    in the label, not closed -- a constrained adjustment is out of scope."""
    res = run_pipeline(collection, RunConfig())
    indices = res["indices"]
    categories = [c for c in indices.columns if c != "All items"][:2]
    cfg = SeasonalConfig(adjustment_engine="stl")
    parts = pd.DataFrame({c: sn.adjust(indices[c], cfg, series_name=c).adjusted
                          for c in categories})
    total = (indices[categories[0]] + indices[categories[1]]) / 2
    direct = sn.adjust(total, cfg, series_name="two categories").adjusted
    indirect = parts.mean(axis=1)
    gap = (direct - indirect.reindex(direct.index)).abs().max()
    assert gap > 1e-6
