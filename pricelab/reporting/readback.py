"""Read the provenance stamp back out of each export format.

Used by the tests that prove every export carries the stamp, and by
anyone handed a file who wants to know which run produced it without
opening the application. Each reader returns the same `ProvenanceStamp`
the exporter embedded, or raises `StampNotFound`.
"""

from __future__ import annotations

import io
import json
import re
import zipfile
from typing import Any

from lxml import etree

from ..core.provenance import STAMP_KEY, ProvenanceStamp


class StampNotFound(ValueError):
    """The file carries no provenance stamp where this format puts it."""


def _parse(text: str) -> ProvenanceStamp:
    try:
        return ProvenanceStamp.from_json(text)
    except (json.JSONDecodeError, TypeError) as exc:
        raise StampNotFound(f"stamp text is not a valid stamp: {exc}") from exc


def from_csv(text: str) -> ProvenanceStamp:
    for line in text.splitlines():
        if line.startswith(f"# {STAMP_KEY} "):
            return _parse(line[len(f"# {STAMP_KEY} "):])
    raise StampNotFound("no stamp comment line")


def from_markdown(text: str) -> ProvenanceStamp:
    match = re.search(rf"```{STAMP_KEY}\n(.*?)\n```", text, re.S)
    if not match:
        raise StampNotFound("no fenced stamp block")
    return _parse(match.group(1))


def from_docx(data: bytes) -> ProvenanceStamp:
    from docx import Document

    doc = Document(io.BytesIO(data))
    for para in doc.paragraphs:
        if para.text.startswith(f"{STAMP_KEY} "):
            return _parse(para.text[len(STAMP_KEY) + 1:])
    raise StampNotFound("no stamp paragraph")


def from_pptx(data: bytes) -> ProvenanceStamp:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    for slide in prs.slides:
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text
            if notes.startswith(f"{STAMP_KEY} "):
                return _parse(notes[len(STAMP_KEY) + 1:])
    raise StampNotFound("no stamp in any slide's notes")


def from_xlsx(data: bytes) -> ProvenanceStamp:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True)
    if "Provenance" not in wb.sheetnames:
        raise StampNotFound("no Provenance sheet")
    for row in wb["Provenance"].iter_rows(values_only=True):
        if row and row[0] == STAMP_KEY and isinstance(row[1], str):
            return _parse(row[1])
    raise StampNotFound("Provenance sheet has no stamp row")


def from_pdf(data: bytes) -> ProvenanceStamp:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    meta: dict[str, Any] = dict(reader.metadata or {})
    keywords = meta.get("/Keywords")
    if keywords:
        try:
            return _parse(str(keywords))
        except StampNotFound:
            pass
    text = "\n".join((page.extract_text() or "") for page in reader.pages)
    match = re.search(rf"{STAMP_KEY}\s+(\{{.*\}})", text.replace("\n", ""), re.S)
    if match:
        return _parse(match.group(1))
    raise StampNotFound("no stamp in the PDF's metadata or text")


def from_sdmx(data: bytes) -> ProvenanceStamp:
    ns = {"common": "http://www.sdmx.org/resources/sdmxml/schemas/v2_1/common"}
    root = etree.fromstring(data)
    for ann in root.iter(f"{{{ns['common']}}}Annotation"):
        title = ann.find("common:AnnotationTitle", ns)
        if title is not None and title.text == STAMP_KEY:
            body = ann.find("common:AnnotationText", ns)
            if body is not None and body.text:
                return _parse(body.text)
    raise StampNotFound("no provenance annotation")


def is_office_zip(data: bytes) -> bool:
    return zipfile.is_zipfile(io.BytesIO(data))
